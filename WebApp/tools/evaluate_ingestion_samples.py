"""Valutazione read-only e anonimizzata dei campioni di ingestion LogiDesk."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

FUNCTIONS_DIR = Path(__file__).resolve().parents[1] / "functions"
sys.path.insert(0, str(FUNCTIONS_DIR))
from services.document_ingestion import extract_ddt  # noqa: E402


def evaluate_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    extracted = [extract_ddt(page.extract_text() or "") for page in reader.pages]
    ddt_pages = [item for item in extracted if item["numero_ddt"]]
    complete = [
        item for item in ddt_pages
        if item["codice_punto_committente"] and item["indirizzo"] and item["codice_zona_logistica"]
    ]
    if ddt_pages:
        status = "READY_FOR_DRY_RUN" if len(complete) == len(ddt_pages) else "OPERATOR_REVIEW_REQUIRED"
        kind = "DDT_MULTI_PAGE"
    else:
        status = "MAPPING_REQUIRED"
        kind = "GIRO_OR_REPORT_PDF"
    return {
        "file": path.name, "kind": kind, "pages": len(reader.pages),
        "ddt_detected": len(ddt_pages), "delivery_points_complete": len(complete),
        "article_occurrences": sum(len(item["articoli"]) for item in ddt_pages),
        "status": status, "writes_performed": 0,
    }


def evaluate_excel(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        names = workbook.sheetnames
        nonempty = sum(
            1 for sheet in workbook.worksheets
            if sum(1 for row in sheet.iter_rows(values_only=True) if any(value not in (None, "") for value in row)) > 1
        )
        is_planning = "Riepilogo" in names and len(names) > 1
        return {
            "file": path.name,
            "kind": "PLANNING_MULTI_SHEET" if is_planning else "DELIVERY_LIST_EXCEL",
            "sheets": len(names), "nonempty_sheets": nonempty,
            "proposed_trips": len(names) - 1 if is_planning else 1,
            "status": "OPERATOR_SHEET_CHOICE_REQUIRED" if is_planning else "MAPPING_REQUIRED",
            "writes_performed": 0,
        }
    finally:
        workbook.close()


def evaluate(path: Path) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return evaluate_pdf(path)
    if suffix in {".xlsx", ".xlsm"}:
        return evaluate_excel(path)
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return {"file": path.name, "kind": "IMAGE_OCR", "status": "OCR_REQUIRED", "writes_performed": 0}
    return {"file": path.name, "kind": "UNSUPPORTED", "status": "REJECTED", "writes_performed": 0}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("folder", type=Path)
    args = parser.parse_args()
    files = sorted(path for path in args.folder.iterdir() if path.is_file())
    results = [evaluate(path) for path in files]
    print(json.dumps({
        "mode": "READ_ONLY_DRY_RUN", "files_total": len(results),
        "writes_performed": 0, "files": results,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
