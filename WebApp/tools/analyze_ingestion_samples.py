"""Inventario read-only di campioni documentali per il collaudo ingestion.

Stampa esclusivamente metadati strutturali. Non copia i file e non mostra il
contenuto delle celle o il testo integrale dei PDF.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader


ANCHORS = (
    "ddt",
    "documento di trasporto",
    "luogo di destinazione",
    "codice destinazione",
    "causale del trasporto",
    "responsabile del trasporto",
    "codice articolo",
    "descrizione",
    "quantita",
    "quantità",
    "data",
    "giro",
    "peso",
    "colli",
)

HEADER_HINTS = (
    "codice", "cliente", "ragione sociale", "nome cliente", "indirizzo",
    "localita", "località", "provincia", "colli", "peso", "quantita",
    "quantità", "giro", "viaggio", "articolo", "descrizione", "data",
    "giorno", "targa", "mezzo", "autista", "ora", "tempo", "fatturato",
    "note", "fascia oraria", "tipo", "resi",
)


def _signature(payload: dict[str, Any]) -> str:
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return "structure-sha256-" + hashlib.sha256(encoded).hexdigest()[:16]


def inspect_pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    pages = []
    all_text = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        normalized = re.sub(r"\s+", " ", text).strip()
        pages.append(
            {
                "page": number,
                "characters": len(normalized),
                "requires_ocr": len(normalized) < 30,
            }
        )
        all_text.append(normalized.lower())
    searchable = " ".join(all_text)
    anchors = sorted(anchor for anchor in ANCHORS if anchor in searchable)
    descriptor = {
        "kind": "pdf",
        "page_count": len(pages),
        "page_text_classes": ["ocr" if p["requires_ocr"] else "digital" for p in pages],
        "anchors": anchors,
    }
    return {
        "kind": "pdf",
        "page_count": len(pages),
        "pages": pages,
        "anchors": anchors,
        "structure_signature": _signature(descriptor),
    }


def inspect_excel(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
            candidates = []
            for index, row in enumerate(matrix[:40]):
                values = [str(value).strip() for value in row if value not in (None, "")]
                if len(values) < 3:
                    continue
                combined = " | ".join(values).lower()
                hints = sum(hint in combined for hint in HEADER_HINTS)
                if hints >= 2:
                    candidates.append({"header_row": index + 1, "header_count": len(values), "hint_count": hints})
            if not candidates:
                ranked = sorted(
                    ((sum(value not in (None, "") for value in row), index) for index, row in enumerate(matrix[:40])),
                    reverse=True,
                )
                if ranked and ranked[0][0] >= 2:
                    candidates.append({"header_row": ranked[0][1] + 1, "header_count": ranked[0][0], "hint_count": 0, "fallback": True})
            sheets.append(
                {
                    "name": worksheet.title,
                    "rows": len(matrix),
                    "columns": max((len(row) for row in matrix), default=0),
                    "table_candidates": candidates,
                    "merged_ranges": len(getattr(getattr(worksheet, "merged_cells", None), "ranges", ())),
                }
            )
    finally:
        workbook.close()
    descriptor = {
        "kind": "excel",
        "sheets": [
            {
                "name": sheet["name"].strip().lower(),
                "columns": sheet["columns"],
                "tables": sheet["table_candidates"],
            }
            for sheet in sheets
        ],
    }
    return {
        "kind": "excel",
        "sheet_count": len(sheets),
        "sheets": sheets,
        "structure_signature": _signature(descriptor),
    }


def inspect(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "file_name": path.name,
        "size": path.stat().st_size,
    }
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        result.update(inspect_pdf(path))
    elif suffix in {".xlsx", ".xlsm"}:
        result.update(inspect_excel(path))
    else:
        result.update({"kind": "unsupported", "error": f"Unsupported extension: {suffix}"})
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("paths", nargs="+")
    args = parser.parse_args()
    results = []
    for raw_path in args.paths:
        path = Path(raw_path)
        if not path.is_file():
            results.append({"file_name": path.name, "kind": "missing", "error": "FILE_NOT_FOUND"})
            continue
        try:
            results.append(inspect(path))
        except Exception as exc:  # report per-file failure without losing the batch
            results.append({"file_name": path.name, "kind": "error", "error": type(exc).__name__})
    print(json.dumps({"files": results}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
